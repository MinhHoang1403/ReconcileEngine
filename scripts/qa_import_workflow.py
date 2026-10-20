"""Automated unfamiliar-file scenarios, not a human comprehension/usability study."""

from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit

import httpx
from openpyxl import Workbook, load_workbook

from scripts.browser_smoke import DevTools, capture, click_button, find_edge, free_port, wait_for_tab, wait_until
from scripts.qa_matching import csv_bytes


def write_source(path: Path, headers: tuple[str, ...], rows: list[tuple]) -> None:
    if path.suffix == ".csv":
        path.write_bytes(csv_bytes(headers, rows))
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        workbook.save(path)
        workbook.close()


def select_value(browser: DevTools, label: str, value: str) -> None:
    assert browser.evaluate(
        "(() => { const label = [...document.querySelectorAll('label')].find(x => "
        "(x.querySelector('.field-label > span') ?? x.querySelector(':scope > span'))?.textContent === "
        f"{json.dumps(label, ensure_ascii=False)}); "
        "const select = label?.querySelector('select'); if (!select) return false; "
        f"select.value = {json.dumps(value)}; select.dispatchEvent(new Event('change', {{bubbles:true}})); "
        "return true; })()"
    )


def exercise(url: str, width: int, height: int, output: Path, cases: list[str]) -> list[dict]:
    address = urlsplit(url)
    if address.scheme != "http" or address.hostname not in {"127.0.0.1", "localhost"} or address.port != 5174:
        raise ValueError("Browser mutation QA is restricted to loopback Lab website port 5174")
    output.mkdir(parents=True, exist_ok=True)
    port = free_port()
    profile = Path(tempfile.mkdtemp(prefix="smartrecon-import-qa-"))
    process = subprocess.Popen(
        [
            str(find_edge()),
            "--headless=new",
            "--disable-gpu",
            "--no-first-run",
            "--no-default-browser-check",
            "--remote-allow-origins=*",
            f"--remote-debugging-port={port}",
            f"--user-data-dir={profile}",
            "about:blank",
        ],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    browser = None
    results = []
    try:
        browser = DevTools(wait_for_tab(port)["webSocketDebuggerUrl"])
        for domain in ("Page", "Runtime", "Network", "Log", "DOM"):
            browser.command(f"{domain}.enable")
        browser.command(
            "Emulation.setDeviceMetricsOverride",
            {
                "width": width,
                "height": height,
                "deviceScaleFactor": 1,
                "mobile": width <= 768,
            },
        )
        with httpx.Client(base_url=url, timeout=90) as client:
            for case in cases:
                count = 1000 if case == "large" else 2
                suffix = ".xlsx" if case == "xlsx" else ".csv"
                orders = [(f"QA-{i:06}", "2026-10-19", str(1000 + i), "VND") for i in range(count)]
                banks = [
                    (f"PAY-{i:06}", "2026-10-20", str(1000 + i), "VND", "CREDIT", "POSTED_CASH", f"QA-{i:06}")
                    for i in range(count)
                ]
                order_headers = ("SO Number", "Created Date", "Amount Due", "CCY")
                bank_headers = ("TX ID", "Posting Date", "Credit Amount", "CCY", "Credit Debit", "Basis", "description")
                if case == "ambiguous":
                    order_headers += ("Balance Due",)
                    orders = [(*row, row[2]) for row in orders]
                if case == "locale":
                    orders = [(row[0], row[1], f"{int(row[2]):,}".replace(",", "."), row[3]) for row in orders]
                    banks = [(row[0], row[1], f"{int(row[2]):,}".replace(",", "."), *row[3:]) for row in banks]
                order_path, bank_path = output / f"{case}-receivables{suffix}", output / f"{case}-payments{suffix}"
                write_source(order_path, order_headers, orders)
                write_source(bank_path, bank_headers, banks)
                browser.command("Page.navigate", {"url": url})
                wait_until(browser, "document.body?.innerText.includes('Hệ thống sẵn sàng')")
                browser.evaluate("""window.qa = {
                        buttons: [], changes: 0, advanced: 0, requests: [], started: performance.now()
                    };
                    const originalFetch = window.fetch;
                    window.fetch = async (...args) => {
                        const response = await originalFetch(...args);
                        const item = {url: String(args[0]), method: args[1]?.method ?? 'GET', status: response.status};
                        try { item.body = await response.clone().json(); } catch { /* Excel/non-JSON */ }
                        window.qa.requests.push(item); return response;
                    };
                    document.addEventListener('click', e => {
                        const b = e.target.closest('button'); if (b) window.qa.buttons.push(b.textContent.trim());
                        if (e.target.closest('summary')) window.qa.advanced += 1;
                    });
                    document.addEventListener('change', () => window.qa.changes += 1);
                """)
                root = browser.command("DOM.getDocument")["root"]["nodeId"]
                for selector, path in (("#orders-file", order_path), ("#bank-file", bank_path)):
                    node = browser.command("DOM.querySelector", {"nodeId": root, "selector": selector})["nodeId"]
                    browser.command("DOM.setFileInputFiles", {"nodeId": node, "files": [str(path.resolve())]})
                click_button(browser, "Kiểm tra và xem trước")
                wait_until(
                    browser,
                    "document.body.innerText.includes('Kiểm tra và chuẩn hóa dữ liệu') || "
                    "document.body.innerText.includes('Dữ liệu sẵn sàng đối soát')",
                    60,
                )
                automatic = browser.evaluate("document.body.innerText.includes('Dữ liệu sẵn sàng đối soát')")
                assert automatic is (case not in {"ambiguous", "locale"}), (
                    case,
                    browser.evaluate("document.body.innerText"),
                )
                if not automatic:
                    capture(browser, output / f"{case}-question-{width}.png")
                    if case == "ambiguous":
                        assert (
                            browser.evaluate(
                                "[...document.querySelectorAll('label')].find(x => "
                                "x.querySelector('.field-label')?.textContent.includes('Số tiền cần thu'))"
                                "?.querySelector('select')?.value"
                            )
                            == ""
                        )
                        select_value(browser, "Số tiền cần thu", "Amount Due")
                    else:
                        assert browser.evaluate("document.body.innerText.includes('hệ thống không tự đoán')")
                        select_value(browser, "Dấu thập phân", ",")
                        select_value(browser, "Dấu phân cách hàng nghìn", ".")
                    click_button(browser, "Xác nhận và chuẩn hóa dữ liệu")
                wait_until(browser, "document.body.innerText.includes('Dữ liệu sẵn sàng đối soát')", 60)
                before_run = browser.evaluate(
                    "({...window.qa, elapsedMilliseconds: performance.now() - window.qa.started})"
                )
                assert not any(x["url"] == "/api/runs" for x in before_run["requests"])
                checked = next(x["body"] for x in before_run["requests"] if x["url"].endswith("/mapping"))
                assert checked["report"]["status"] == "READY"
                assert checked["acknowledged_issue_codes"] == []
                capture(browser, output / f"{case}-health-{width}.png")
                expert_path = False
                if case == "csv":
                    assert browser.evaluate(
                        "[...document.querySelectorAll('summary')].find(x => "
                        "x.textContent === 'Xem file gốc và cách đọc dữ liệu')?.click(); true"
                    )
                    wait_until(browser, "document.querySelectorAll('.normalized-card tbody tr').length === 4")
                    assert browser.evaluate("document.documentElement.scrollWidth") <= width
                    click_button(browser, "Chỉnh cách đọc dữ liệu")
                    wait_until(browser, "document.body.innerText.includes('Kiểm tra và chuẩn hóa dữ liệu')")
                    browser.evaluate(
                        "[...document.querySelectorAll('summary')].find(x => "
                        "x.textContent === 'Xem trước dữ liệu sau chuẩn hóa')?.click()"
                    )
                    wait_until(browser, "document.querySelectorAll('.normalized-card tbody tr').length === 4")
                    click_button(browser, "Xác nhận và chuẩn hóa dữ liệu")
                    wait_until(browser, "document.body.innerText.includes('Dữ liệu sẵn sàng đối soát')")
                    assert browser.evaluate("document.documentElement.scrollWidth") <= width
                    expert_path = True
                click_button(browser, "Chạy đối soát")
                wait_until(browser, "document.body.innerText.includes('Kết quả đối soát')", 90)
                layout = browser.evaluate(
                    "({width: innerWidth, scroll: document.documentElement.scrollWidth, "
                    "overflowing: [...document.querySelectorAll('body *')]"
                    ".filter(x => x.getBoundingClientRect().right > "
                    + str(width + 1)
                    + " && x.getBoundingClientRect().width > 0).slice(0,15)"
                    ".map(x => ({tag:x.tagName, cls:x.className, right:x.getBoundingClientRect().right}))})"
                )
                if layout["scroll"] > width:
                    layout["containers"] = browser.evaluate(
                        "[...document.querySelectorAll('main,section,article,div,details,summary,svg')]"
                        ".filter(x => x.getBoundingClientRect().right > "
                        + str(width + 1)
                        + " && !x.closest('table')).map(x => ({tag:x.tagName, cls:x.getAttribute('class'), "
                        "width:x.getBoundingClientRect().width, right:x.getBoundingClientRect().right, "
                        "overflow:getComputedStyle(x).overflowX})).slice(0,35)"
                    )
                assert layout["scroll"] <= width, layout
                capture(browser, output / f"{case}-dashboard-{width}.png")
                href = browser.evaluate(
                    "document.querySelector('a[href*=\"exports?format=xlsx\"]')?.getAttribute('href')"
                )
                run_id = href.split("/")[3]
                response = client.get(f"/api/runs/{run_id}/dashboard")
                response.raise_for_status()
                dashboard = response.json()
                cash = sum(1000 + i for i in range(count))
                assert dashboard["conservation"]["is_valid"]
                assert dashboard["kpis"]["confirmed_cash_allocated_minor"] == cash
                assert dashboard["kpis"]["unapplied_bank_credit_minor"] == 0
                assert dashboard["payment_status_counts"]["FULL"] == count
                requests = browser.evaluate("window.qa.requests")
                order_page = next(x["body"] for x in requests if "/orders?" in x["url"])
                assert len(order_page["items"]) == min(count, 50)
                exported = client.get(href)
                exported.raise_for_status()
                workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
                assert len(workbook.sheetnames) == 9
                workbook.close()
                assert not browser.failures, browser.failures
                results.append(
                    {
                        "case": case,
                        "rows": count,
                        "viewport": [width, height],
                        "run_id": run_id,
                        "automatic": automatic,
                        "clicks_to_health": len(before_run["buttons"]),
                        "buttons": before_run["buttons"],
                        "changes_including_files": before_run["changes"],
                        "advanced_openings": before_run["advanced"],
                        "elapsed_ms": before_run["elapsedMilliseconds"],
                        "cash_minor": cash,
                        "conservation": True,
                        "initial_order_page_size": len(order_page["items"]),
                        "export_bytes": len(exported.content),
                        "source_sha256": {
                            path.name: hashlib.sha256(path.read_bytes()).hexdigest() for path in (order_path, bank_path)
                        },
                        "layout": layout,
                        "human_usability_tested": False,
                        "expert_preview_back_reconfirm_verified": expert_path,
                    }
                )
                (output / f"results-{width}.json").write_text(
                    json.dumps(results, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
                )
        return results
    finally:
        if browser:
            browser.close()
        process.terminate()
        try:
            process.wait(timeout=5)
        except subprocess.TimeoutExpired:
            process.kill()
        shutil.rmtree(profile, ignore_errors=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default="http://127.0.0.1:5174")
    parser.add_argument("--width", type=int, default=1440)
    parser.add_argument("--height", type=int, default=900)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument(
        "--cases",
        nargs="+",
        choices=["csv", "xlsx", "ambiguous", "locale", "large"],
        default=["csv", "xlsx", "ambiguous", "locale", "large"],
    )
    args = parser.parse_args()
    print(json.dumps(exercise(args.url, args.width, args.height, args.output, args.cases), ensure_ascii=False))


if __name__ == "__main__":
    main()
