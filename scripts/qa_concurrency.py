"""Exercise review concurrency and revision invariants through independent HTTP requests."""

from __future__ import annotations

import argparse
import json
import time
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from itertools import combinations
from pathlib import Path
from threading import Barrier
from typing import Any

from openpyxl import load_workbook


class ApiClient:
    def __init__(self, base_url: str) -> None:
        self.base_url = base_url.rstrip("/")

    def request(self, method: str, path: str, payload: dict[str, Any] | None = None) -> tuple[int, Any, float]:
        body = json.dumps(payload).encode() if payload is not None else None
        request = urllib.request.Request(
            self.base_url + path,
            data=body,
            method=method,
            headers={"Content-Type": "application/json"} if body is not None else {},
        )
        started = time.perf_counter()
        try:
            with urllib.request.urlopen(request, timeout=60) as response:
                content = response.read()
                elapsed = time.perf_counter() - started
                if "application/json" in response.headers.get("Content-Type", ""):
                    return response.status, json.loads(content), elapsed
                return response.status, content, elapsed
        except urllib.error.HTTPError as error:
            content = error.read()
            elapsed = time.perf_counter() - started
            try:
                detail = json.loads(content)
            except json.JSONDecodeError:
                detail = {"body": content.decode(errors="replace")}
            return error.code, detail, elapsed

    def json(self, method: str, path: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
        status, content, _ = self.request(method, path, payload)
        if status != 200 or not isinstance(content, dict):
            raise AssertionError(f"{method} {path} returned {status}: {content}")
        return content


def _new_demo_run(client: ApiClient) -> str:
    inspection = client.json("POST", "/api/demo/datasets")
    mapping = {
        **inspection["suggested_mapping"],
        "currency": "VND",
        "source_profile": "BANK_CREDIT",
        "date_format": "%Y-%m-%d",
        "decimal_separator": ".",
        "thousand_separator": ",",
        "order_defaults": {"order_status": "ELIGIBLE"},
        "bank_defaults": {},
    }
    gate = client.json("POST", f"/api/uploads/{inspection['session_id']}/mapping", mapping)
    warning_codes = [item["code"] for item in gate["report"]["issues"] if item["severity"] == "WARNING"]
    client.json(
        "POST",
        f"/api/uploads/{inspection['session_id']}/resolutions",
        {"exclusions": [], "acknowledged_issue_codes": warning_codes},
    )
    created = client.json(
        "POST",
        "/api/runs",
        {"upload_session_id": inspection["session_id"], "policy": {"currency": "VND"}},
    )
    return str(created["run_id"])


def _competing_pair(items: list[dict[str, Any]]) -> tuple[dict[str, Any], dict[str, Any]]:
    for first, second in combinations(items, 2):
        if first["can_approve"] and second["can_approve"] and set(first["bank_txn_ids"]) & set(second["bank_txn_ids"]):
            return first, second
    raise AssertionError("The demo did not expose two independently approvable competing proposals")


def _command(group_id: str, key: str, note: str, revision: int) -> dict[str, Any]:
    return {
        "group_id": group_id,
        "expected_revision": revision,
        "idempotency_key": key,
        "actor_label": "qa-concurrency",
        "note": note,
    }


def _assert_revision_invariants(client: ApiClient, run_id: str, revision: int) -> dict[str, Any]:
    dashboard = client.json("GET", f"/api/runs/{run_id}/dashboard?revision={revision}")
    orders = client.json("GET", f"/api/runs/{run_id}/orders?revision={revision}&page_size=200")["items"]
    banks = client.json("GET", f"/api/runs/{run_id}/bank-transactions?revision={revision}&page_size=200")["items"]
    kpis = dashboard["kpis"]
    assert dashboard["conservation"]["is_valid"] is True
    assert dashboard["conservation"]["violations"] == []

    inferred_adjustments = 0
    for order in orders:
        adjustment = order["expected_amount_minor"] - order["allocated_cash_minor"] - order["outstanding_minor"]
        assert adjustment >= 0
        inferred_adjustments += adjustment
    eligible_banks = [bank for bank in banks if bank["unapplied_minor"] is not None]
    for bank in eligible_banks:
        assert bank["posted_amount_minor"] == bank["allocated_cash_minor"] + bank["unapplied_minor"]

    assert sum(order["expected_amount_minor"] for order in orders) == kpis["expected_receivable_minor"]
    assert sum(order["allocated_cash_minor"] for order in orders) == kpis["confirmed_cash_allocated_minor"]
    assert sum(order["outstanding_minor"] for order in orders) == kpis["outstanding_receivable_minor"]
    assert inferred_adjustments == kpis["verified_adjustments_minor"]
    assert sum(bank["allocated_cash_minor"] for bank in eligible_banks) == kpis["confirmed_cash_allocated_minor"]
    assert sum(bank["unapplied_minor"] for bank in eligible_banks) == kpis["unapplied_bank_credit_minor"]
    return dashboard


def _assert_excel_matches_dashboard(workbook_bytes: bytes, dashboard: dict[str, Any]) -> None:
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=False, read_only=True)
    assert len(workbook.sheetnames) == 9
    allocation = workbook["05_AllocationEdges"]
    headers = [cell.value for cell in next(allocation.iter_rows(min_row=1, max_row=1))]
    amount_index = headers.index("Số tiền phân bổ")
    total = sum(row[amount_index].value for row in allocation.iter_rows(min_row=2))
    assert total == dashboard["kpis"]["confirmed_cash_allocated_minor"]
    assert all(cell.data_type != "f" for sheet in workbook.worksheets for row in sheet.iter_rows() for cell in row)
    workbook.close()


def run(base_url: str) -> dict[str, Any]:
    client = ApiClient(base_url)
    run_id = _new_demo_run(client)
    queue = client.json("GET", f"/api/runs/{run_id}/reviews?page_size=200")
    first, second = _competing_pair(queue["items"])
    payloads = [
        _command(first["group_id"], "qa-concurrent-approve-a", "Concurrent approval A", 0),
        _command(second["group_id"], "qa-concurrent-approve-b", "Concurrent approval B", 0),
    ]

    start_barrier = Barrier(3)

    def concurrent_approve(payload: dict[str, Any]) -> tuple[int, Any, float]:
        start_barrier.wait()
        return client.request("POST", f"/api/runs/{run_id}/reviews/approve", payload)

    with ThreadPoolExecutor(max_workers=2) as executor:
        futures = [executor.submit(concurrent_approve, item) for item in payloads]
        start_barrier.wait()
        concurrent_results = [future.result() for future in futures]
    assert sorted(item[0] for item in concurrent_results) == [200, 409]
    success_index = next(index for index, item in enumerate(concurrent_results) if item[0] == 200)
    conflict = next(item[1] for item in concurrent_results if item[0] == 409)
    assert conflict["detail"]["error_code"] == "STALE_REVISION"
    successful_payload = payloads[success_index]
    successful_response = concurrent_results[success_index][1]
    assert successful_response["effective_revision"] == 1
    assert successful_response["conservation"]["is_valid"] is True

    replay_status, replay, _ = client.request("POST", f"/api/runs/{run_id}/reviews/approve", successful_payload)
    assert replay_status == 200 and replay["idempotent_replay"] is True
    changed_payload = {**successful_payload, "note": "A different command with the same key"}
    reused_status, reused, _ = client.request("POST", f"/api/runs/{run_id}/reviews/approve", changed_payload)
    assert reused_status == 409 and reused["detail"]["error_code"] == "IDEMPOTENCY_KEY_REUSED"

    many_to_many = next(item for item in queue["items"] if len(item["order_ids"]) > 1 and len(item["bank_txn_ids"]) > 1)
    rollback_status, rollback, _ = client.request(
        "POST",
        f"/api/runs/{run_id}/reviews/approve",
        _command(many_to_many["group_id"], "qa-rollback-many-to-many", "Must be rejected", 1),
    )
    assert rollback_status == 409 and rollback["detail"]["error_code"] == "COMPLEX_GROUP_NOT_APPROVABLE"

    losing_payload = payloads[1 - success_index]
    rejected = client.json(
        "POST",
        f"/api/runs/{run_id}/reviews/reject",
        _command(losing_payload["group_id"], "qa-reject-competitor", "Reject competing proposal", 1),
    )
    assert rejected["effective_revision"] == 2 and rejected["allocation_delta_minor"] == 0
    dependency_status, dependency, _ = client.request(
        "POST",
        f"/api/runs/{run_id}/reviews/undo",
        {
            "action_id": successful_response["action_id"],
            "expected_revision": 2,
            "idempotency_key": "qa-undo-dependent-approval",
            "actor_label": "qa-concurrency",
            "note": "Must not erase the downstream rejection",
        },
    )
    assert dependency_status == 409
    assert dependency["detail"]["error_code"] == "UNDO_DEPENDENCY_CONFLICT"
    undone = client.json(
        "POST",
        f"/api/runs/{run_id}/reviews/undo",
        {
            "action_id": rejected["action_id"],
            "expected_revision": 2,
            "idempotency_key": "qa-undo-rejection",
            "actor_label": "qa-concurrency",
            "note": "Restore the rejected proposal",
        },
    )
    assert undone["effective_revision"] == 3 and undone["conservation"]["is_valid"] is True

    stale_status, stale, _ = client.request(
        "POST",
        f"/api/runs/{run_id}/reviews/reject",
        _command(first["group_id"], "qa-explicit-stale", "Expected stale revision", 0),
    )
    assert stale_status == 409 and stale["detail"]["error_code"] == "STALE_REVISION"

    dashboards = [_assert_revision_invariants(client, run_id, revision) for revision in range(4)]
    history = client.json("GET", f"/api/runs/{run_id}/reviews/history?page_size=200")
    assert history["total"] == 3
    assert {item["action_type"] for item in history["items"]} == {
        "APPROVE_GROUP",
        "REJECT_GROUP",
        "UNDO_ACTION",
    }
    run = client.json("GET", f"/api/runs/{run_id}")
    assert run["effective_revision"] == 3

    export_status, workbook_bytes, export_seconds = client.request(
        "GET", f"/api/runs/{run_id}/exports?format=xlsx&expected_revision=3"
    )
    assert export_status == 200 and isinstance(workbook_bytes, bytes)
    dashboard_started = time.perf_counter()
    current_dashboard = client.json("GET", f"/api/runs/{run_id}/dashboard?revision=3")
    dashboard_seconds = time.perf_counter() - dashboard_started
    _assert_excel_matches_dashboard(workbook_bytes, current_dashboard)

    return {
        "run_id": run_id,
        "concurrent_statuses": [item[0] for item in concurrent_results],
        "concurrent_request_seconds": [item[2] for item in concurrent_results],
        "committed_revision": 3,
        "audit_actions": history["total"],
        "revision_conservation": [item["conservation"]["is_valid"] for item in dashboards],
        "dashboard_seconds_after_review": dashboard_seconds,
        "export_seconds_after_review": export_seconds,
        "export_bytes": len(workbook_bytes),
        "idempotent_replay": replay["idempotent_replay"],
        "rollback_error": rollback["detail"]["error_code"],
        "undo_dependency_error": dependency["detail"]["error_code"],
        "stale_error": stale["detail"]["error_code"],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--url", default="http://127.0.0.1:8000")
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = run(args.url)
    encoded = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(encoded + "\n", encoding="utf-8")
    print(encoded)


if __name__ == "__main__":
    main()
