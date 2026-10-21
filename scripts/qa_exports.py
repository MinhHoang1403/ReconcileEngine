"""Verify simultaneous first downloads through PostgreSQL, restricted to the Lab API."""

import argparse
import hashlib
import json
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO
from pathlib import Path
from threading import Barrier
from urllib.parse import urlsplit

from openpyxl import load_workbook

from scripts.qa_artifacts import audit
from scripts.qa_concurrency import ApiClient, _new_demo_run


def exercise(base_url: str, directory: Path) -> dict:
    address = urlsplit(base_url)
    if address.scheme != "http" or address.hostname not in {"localhost", "127.0.0.1"} or address.port != 8002:
        raise ValueError("Mutation QA is restricted to loopback Lab API port 8002")
    client = ApiClient(base_url)
    run_id = _new_demo_run(client)
    before = audit("smartrecon_lab", directory)
    barrier = Barrier(6)

    def download(_: int) -> bytes:
        barrier.wait(timeout=15)
        status, content, _ = ApiClient(base_url).request(
            "GET", f"/api/runs/{run_id}/exports?format=xlsx&expected_revision=0"
        )
        assert status == 200
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
        assert len(workbook.sheetnames) == 9
        workbook.close()
        return content

    with ThreadPoolExecutor(max_workers=6) as executor:
        contents = list(executor.map(download, range(6)))
    hashes = {hashlib.sha256(content).hexdigest() for content in contents}
    assert len(hashes) == 1
    after = audit("smartrecon_lab", directory)
    assert not after["mismatches"], after["mismatches"]
    assert after["total"] == before["total"] + 1
    summary = client.json("GET", f"/api/runs/{run_id}/dashboard?revision=0")
    assert summary["conservation"]["is_valid"]
    assert summary["kpis"]["confirmed_cash_allocated_minor"] == 66431000
    return {
        "run_id": run_id,
        "simultaneous_downloads": 6,
        "all_http_200": True,
        "distinct_sha256": sorted(hashes),
        "export_bytes": len(contents[0]),
        "manifest_rows_added": after["total"] - before["total"],
        "artifact_mismatches": 0,
        "cash_minor": 66431000,
        "conservation": True,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--url", default="http://127.0.0.1:8002")
    parser.add_argument("--directory", type=Path, default=Path(__file__).resolve().parents[1])
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = exercise(args.url, args.directory.resolve())
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result))


if __name__ == "__main__":
    main()
