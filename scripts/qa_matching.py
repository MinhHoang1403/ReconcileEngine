"""Exercise synthetic financial safety oracles through the isolated Lab HTTP workflow."""

from __future__ import annotations

import argparse
import csv
import json
from io import BytesIO, StringIO
from pathlib import Path
from urllib.parse import urlsplit

import httpx
from openpyxl import load_workbook


def csv_bytes(headers: tuple[str, ...], rows: list[tuple]) -> bytes:
    stream = StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return stream.getvalue().encode("utf-8")


def exercise(base_url: str) -> list[dict]:
    address = urlsplit(base_url)
    if address.scheme != "http" or address.hostname not in {"127.0.0.1", "localhost"} or address.port != 8002:
        raise ValueError("Mutation QA is restricted to the established loopback Lab API on port 8002")
    day = "2026-10-20"
    orders = [("ORDER-A", day, 100000, "SHARED-PAYOUT"), ("ORDER-B", day, 200000, "SHARED-PAYOUT")]
    plain = [(identifier, when, amount, "") for identifier, when, amount, _ in orders]

    def bank(identifier: str, amount: int, memo: str) -> tuple:
        return (identifier, day, amount, "CREDIT", "POSTED_CASH", memo)

    cases = [
        ("reference_collision", orders, [bank("B1", 100000, "ORDER-A SHARED-PAYOUT")], 0),
        (
            "collision_propagates_to_split",
            orders,
            [bank("B1", 50000, "ORDER-A SHARED-PAYOUT"), bank("B2", 50000, "ORDER-A")],
            0,
        ),
        ("unicode_token_boundary", [("DH01", day, 100000, "")], [bank("B1", 100000, "ĐDH01")], 0),
        ("amount_date_is_not_identity", [("ORDER-A", day, 100000, "")], [bank("B1", 100000, "Thanh toan")], 0),
        ("full_reference_group_control", plain, [bank("B1", 300000, "ORDER-A ORDER-B")], 300000),
        (
            "complex_graph_is_review_only",
            plain,
            [bank("B1", 150000, "ORDER-A ORDER-B"), bank("B2", 150000, "ORDER-B")],
            0,
        ),
    ]
    output = []
    with httpx.Client(base_url=base_url.rstrip("/"), timeout=60) as client:

        def post(path: str, payload: dict) -> dict:
            response = client.post(path, json=payload)
            response.raise_for_status()
            return response.json()

        for name, order_rows, bank_rows, expected_cash in cases:
            response = client.post(
                "/api/uploads/inspect",
                files={
                    "orders_file": (
                        "orders.csv",
                        csv_bytes(("order_id", "ordered_at", "expected_amount", "external_ref"), order_rows),
                        "text/csv",
                    ),
                    "bank_file": (
                        "bank.csv",
                        csv_bytes(
                            ("bank_txn_id", "txn_date", "posted_amount", "direction", "amount_basis", "description"),
                            bank_rows,
                        ),
                        "text/csv",
                    ),
                },
            )
            response.raise_for_status()
            session_id = response.json()["session_id"]
            mapping = {
                "order_columns": {name: name for name in ("order_id", "ordered_at", "expected_amount", "external_ref")},
                "bank_columns": {
                    name: name
                    for name in ("bank_txn_id", "txn_date", "posted_amount", "direction", "amount_basis", "description")
                },
                "currency": "VND",
                "date_format": "%Y-%m-%d",
                "confirmed_structure": True,
            }
            health = post(f"/api/uploads/{session_id}/mapping", mapping)
            assert health["report"]["status"] == "READY"
            warnings = [issue["code"] for issue in health["report"]["issues"] if issue["severity"] == "WARNING"]
            if warnings:
                post(f"/api/uploads/{session_id}/resolutions", {"exclusions": [], "acknowledged_issue_codes": warnings})
            run = post("/api/runs", {"upload_session_id": session_id, "policy": {"currency": "VND"}})
            run_id = run["run_id"]
            response = client.get(f"/api/runs/{run_id}/dashboard")
            response.raise_for_status()
            dashboard = response.json()
            assert dashboard["conservation"]["is_valid"]
            assert dashboard["kpis"]["confirmed_cash_allocated_minor"] == expected_cash
            assert dashboard["kpis"]["unapplied_bank_credit_minor"] == sum(row[2] for row in bank_rows) - expected_cash
            exported = client.get(f"/api/runs/{run_id}/exports?format=xlsx&expected_revision=0")
            exported.raise_for_status()
            workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
            assert len(workbook.sheetnames) == 9
            workbook.close()
            output.append(
                {
                    "case": name,
                    "run_id": run_id,
                    "cash_minor": expected_cash,
                    "conservation": True,
                    "export_bytes": len(exported.content),
                }
            )
    return output


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--base-url", default="http://127.0.0.1:8002")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = exercise(args.base_url)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result))


if __name__ == "__main__":
    main()
