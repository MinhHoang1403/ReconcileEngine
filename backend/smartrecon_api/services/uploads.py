"""Upload persistence, inspection, mapping, and Health Gate orchestration."""

from __future__ import annotations

import csv
import shutil
from dataclasses import asdict
from pathlib import Path
from typing import Any
from uuid import uuid4

from fastapi import UploadFile
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.smartrecon_api.db.models import ArtifactManifest, UploadSession
from backend.smartrecon_api.schemas.uploads import MappingRequest, ResolutionRequest
from backend.smartrecon_api.services.artifacts import (
    relative_artifact_path,
    resolve_artifact_path,
    upload_directory,
)
from backend.smartrecon_api.services.serialization import to_jsonable
from smartrecon.ingest.eda import eda_to_dict, summarize_raw, summarize_standardized
from smartrecon.ingest.health import Exclusion, apply_exclusions, build_health_report
from smartrecon.ingest.mapping import (
    BANK_REQUIRED,
    ORDER_REQUIRED,
    MappingConfig,
    default_demo_mapping,
    suggest_mapping,
)
from smartrecon.ingest.normalization import NormalizedBundle, normalize_bundle
from smartrecon.ingest.readers import MAX_FILE_BYTES, RawSourceBundle, read_source_bundle

ALLOWED_CONTENT_TYPES = {
    "text/csv",
    "application/csv",
    "application/vnd.ms-excel",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",
}


async def save_upload(upload: UploadFile, directory: Path, target_name: str) -> Path:
    """Stream an upload to a server-chosen path with a hard byte limit."""
    if not upload.filename:
        raise ValueError("Upload thiếu tên file")
    suffix = Path(upload.filename).suffix.lower()
    if suffix not in {".csv", ".xlsx"}:
        raise ValueError("Chỉ hỗ trợ CSV hoặc XLSX")
    if upload.content_type and upload.content_type not in ALLOWED_CONTENT_TYPES:
        raise ValueError(f"Content-Type không được hỗ trợ: {upload.content_type}")
    path = directory / f"{target_name}{suffix}"
    byte_count = 0
    with path.open("wb") as handle:
        while chunk := await upload.read(1024 * 1024):
            byte_count += len(chunk)
            if byte_count > MAX_FILE_BYTES:
                handle.close()
                path.unlink(missing_ok=True)
                raise ValueError("File vượt giới hạn 25 MB")
            handle.write(chunk)
    await upload.close()
    return path


def create_demo_session(db: Session) -> dict[str, Any]:
    source = Path("data/demo/vn_synthetic_demo")
    if not source.is_dir():
        raise FileNotFoundError("Không tìm thấy dữ liệu demo đã sinh")
    session_id = str(uuid4())
    directory = upload_directory(session_id)
    orders_path = directory / "orders.csv"
    bank_path = directory / "bank.csv"
    shutil.copy2(source / "internal_receivables.csv", orders_path)
    shutil.copy2(source / "bank_transactions.csv", bank_path)
    return _persist_inspection(
        db,
        session_id,
        read_source_bundle(orders_path=orders_path, bank_path=bank_path),
        {"mode": "separate", "orders": str(orders_path), "bank": str(bank_path)},
        mapping=default_demo_mapping(),
        dataset_profile="vn_synthetic_demo",
    )


def create_upload_session(
    db: Session,
    session_id: str,
    raw: RawSourceBundle,
    source_paths: dict[str, Any],
) -> dict[str, Any]:
    return _persist_inspection(db, session_id, raw, source_paths)


def _persist_inspection(
    db: Session,
    session_id: str,
    raw: RawSourceBundle,
    source_paths: dict[str, Any],
    *,
    mapping: MappingConfig | None = None,
    dataset_profile: str = "uploaded",
) -> dict[str, Any]:
    order_suggestions = suggest_mapping(
        list(raw.orders.frame.columns),
        ORDER_REQUIRED
        | {"currency", "order_status", "external_ref", "customer_name", "payment_channel", "description"},
    )
    bank_suggestions = suggest_mapping(
        list(raw.bank_transactions.frame.columns),
        BANK_REQUIRED
        | {"currency", "txn_ref", "value_date", "description", "sender_name", "channel", "batch_ref", "declared_fee"},
    )
    if mapping:
        order_suggestions = mapping.order_columns
        bank_suggestions = mapping.bank_columns
    manifest = {
        "dataset_profile": dataset_profile,
        "source_paths": {
            key: relative_artifact_path(Path(value)) if key in {"orders", "bank", "workbook"} else value
            for key, value in source_paths.items()
        },
        "files": {
            "orders": {
                "name": raw.orders.original_filename,
                "sha256": raw.orders.file_sha256,
                "sheet": raw.orders.sheet_name,
            },
            "bank": {
                "name": raw.bank_transactions.original_filename,
                "sha256": raw.bank_transactions.file_sha256,
                "sheet": raw.bank_transactions.sheet_name,
            },
        },
    }
    session = UploadSession(id=session_id, status="INSPECTED", source_manifest=manifest)
    db.add(session)
    # Explicitly establish the parent row before artifact FKs; the models do not
    # carry ORM relationships because snapshots are queried by their identifiers.
    db.flush()
    for feed, table, path_key in (
        ("RAW_ORDERS", raw.orders, "orders"),
        ("RAW_BANK", raw.bank_transactions, "bank"),
    ):
        relative = manifest["source_paths"].get(path_key) or manifest["source_paths"].get("workbook")
        db.add(
            ArtifactManifest(
                upload_session_id=session_id,
                artifact_type=feed,
                relative_path=relative,
                sha256=table.file_sha256,
                byte_size=resolve_artifact_path(relative).stat().st_size,
            )
        )
    db.commit()
    return {
        "session_id": session_id,
        "status": session.status,
        "dataset_profile": dataset_profile,
        "raw_eda": {
            "orders": eda_to_dict(summarize_raw(raw.orders)),
            "bank_transactions": eda_to_dict(summarize_raw(raw.bank_transactions)),
        },
        "suggested_mapping": {
            "order_columns": order_suggestions,
            "bank_columns": bank_suggestions,
        },
    }


def mapping_from_request(request: MappingRequest | dict[str, Any]) -> MappingConfig:
    validated = request if isinstance(request, MappingRequest) else MappingRequest.model_validate(request)
    payload = validated.model_dump()
    return MappingConfig(**payload)


def load_raw_bundle(session: UploadSession) -> RawSourceBundle:
    paths = session.source_manifest["source_paths"]
    if paths["mode"] == "workbook":
        return read_source_bundle(
            workbook_path=resolve_artifact_path(paths["workbook"]),
            orders_sheet=paths["orders_sheet"],
            bank_sheet=paths["bank_sheet"],
        )
    return read_source_bundle(
        orders_path=resolve_artifact_path(paths["orders"]),
        bank_path=resolve_artifact_path(paths["bank"]),
    )


def read_source_rows(
    session: UploadSession,
    feed: str,
    row_numbers: list[int],
    *,
    limit: int = 50,
) -> dict[int, dict[str, Any]]:
    """Stream selected source rows without accepting a client-controlled artifact path."""
    if feed not in {"orders", "bank_transactions"}:
        raise ValueError("Nguồn dữ liệu không hợp lệ")
    targets = sorted({number for number in row_numbers if number >= 2})
    if len(targets) > limit:
        raise ValueError(f"Chỉ được đọc tối đa {limit} dòng nguồn mỗi yêu cầu")
    if not targets:
        return {}
    paths = session.source_manifest["source_paths"]
    path_key = "orders" if feed == "orders" else "bank"
    relative = paths.get(path_key) or paths.get("workbook")
    path = resolve_artifact_path(relative)
    if path.suffix.lower() == ".csv":
        return _stream_csv_rows(path, targets)
    file_key = "orders" if feed == "orders" else "bank"
    sheet_name = session.source_manifest.get("files", {}).get(file_key, {}).get("sheet")
    return _stream_xlsx_rows(path, sheet_name, targets)


def _stream_csv_rows(path: Path, targets: list[int]) -> dict[int, dict[str, Any]]:
    found: dict[int, dict[str, Any]] = {}
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row_number, row in enumerate(reader, start=2):
            if row_number in targets:
                found[row_number] = {str(key): value for key, value in row.items()}
            if row_number >= targets[-1]:
                break
    return found


def _stream_xlsx_rows(path: Path, sheet_name: str | None, targets: list[int]) -> dict[int, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        headers = [str(value or "") for value in next(iterator)]
        found: dict[int, dict[str, Any]] = {}
        for row_number, values in enumerate(iterator, start=2):
            if row_number in targets:
                found[row_number] = {
                    header: "" if value is None else str(value) for header, value in zip(headers, values, strict=False)
                }
            if row_number >= targets[-1]:
                break
        return found
    finally:
        workbook.close()


def confirm_mapping(db: Session, session: UploadSession, request: MappingRequest) -> dict[str, Any]:
    mapping = mapping_from_request(request)
    raw = load_raw_bundle(session)
    normalized = normalize_bundle(raw, mapping)
    report = build_health_report(normalized)
    standardized = summarize_standardized(
        list(normalized.orders), list(normalized.bank_transactions), mapping.currency.value
    )
    session.mapping_snapshot = to_jsonable(mapping)
    session.health_snapshot = {
        "report": to_jsonable(report),
        "exclusions": [],
        "acknowledged_issue_codes": [],
        "standardized_eda": eda_to_dict(standardized),
        "issue_details": _health_issue_details(raw, normalized, report, session.source_manifest),
    }
    session.status = "HEALTH_READY" if report.status == "READY" else "HEALTH_BLOCKED"
    db.commit()
    return session.health_snapshot | {"session_id": session.id, "status": session.status}


def normalize_session(session: UploadSession) -> tuple[NormalizedBundle, Any]:
    if not session.mapping_snapshot:
        raise ValueError("Mapping chưa được xác nhận")
    normalized = normalize_bundle(load_raw_bundle(session), mapping_from_request(session.mapping_snapshot))
    exclusions = [Exclusion(**item) for item in (session.health_snapshot or {}).get("exclusions", [])]
    if exclusions:
        return apply_exclusions(normalized, exclusions)
    return normalized, build_health_report(normalized)


def apply_resolution(db: Session, session: UploadSession, request: ResolutionRequest) -> dict[str, Any]:
    raw = load_raw_bundle(session)
    normalized = normalize_bundle(raw, mapping_from_request(session.mapping_snapshot or {}))
    row_limits = {"orders": len(raw.orders.frame) + 1, "bank_transactions": len(raw.bank_transactions.frame) + 1}
    for exclusion in request.exclusions:
        if exclusion.source_row_number > row_limits[exclusion.feed]:
            raise ValueError(f"Dòng {exclusion.source_row_number} không tồn tại trong nguồn {exclusion.feed}")
    exclusions = [Exclusion(**item.model_dump()) for item in request.exclusions]
    resolved, report = apply_exclusions(normalized, exclusions)
    mapping = mapping_from_request(session.mapping_snapshot or {})
    standardized = summarize_standardized(
        list(resolved.orders), list(resolved.bank_transactions), mapping.currency.value
    )
    session.health_snapshot = {
        "report": to_jsonable(report),
        "exclusions": to_jsonable(exclusions),
        "acknowledged_issue_codes": sorted(set(request.acknowledged_issue_codes)),
        "standardized_eda": eda_to_dict(standardized),
        "issue_details": _health_issue_details(raw, resolved, report, session.source_manifest),
    }
    session.status = "HEALTH_READY" if report.status == "READY" else "HEALTH_BLOCKED"
    db.commit()
    return session.health_snapshot | {"session_id": session.id, "status": session.status}


def health_issue_detail(session: UploadSession, issue_index: int) -> dict[str, Any]:
    """Return a bounded detail snapshot produced during mapping, never an arbitrary file path."""
    details = (session.health_snapshot or {}).get("issue_details", [])
    if issue_index < 0 or issue_index >= len(details):
        raise LookupError("HEALTH_ISSUE_NOT_FOUND")
    return details[issue_index]


def _health_issue_details(
    raw: RawSourceBundle,
    normalized: NormalizedBundle,
    report: Any,
    manifest: dict[str, Any],
) -> list[dict[str, Any]]:
    """Persist at most 50 related rows per issue while the full mapping pass is already in memory."""
    canonical_by_row: dict[tuple[str, int], dict[str, Any]] = {}
    for feed, records in (("orders", normalized.orders), ("bank_transactions", normalized.bank_transactions)):
        for record in records:
            if record.lineage:
                canonical_by_row[(feed, record.lineage.source_row_number)] = to_jsonable(asdict(record))

    file_meta = manifest.get("files", {})
    tables = {"orders": raw.orders, "bank_transactions": raw.bank_transactions}
    actions = {
        "POSSIBLE_REIMPORT": (
            "So sánh hai dòng nguồn và mã giao dịch; chỉ loại trừ khi xác minh đây là cùng một giao dịch nhập lặp."
        ),
        "MISSING_REFERENCE_EVIDENCE": (
            "Bổ sung mã tham chiếu hoặc nội dung giao dịch nếu nguồn có dữ liệu; "
            "nếu không, giao dịch chỉ được đề xuất ghép."
        ),
        "ORDER_ID_COLLISION": "Sửa mã đơn hàng hoặc cấu hình khóa ghép có thể phân biệt từng nghĩa vụ.",
        "BANK_ID_COLLISION": "Sửa mã giao dịch bị trùng trước khi chạy đối soát.",
        "ORDER_PARSE_ERROR": "Sửa giá trị sai định dạng trong file nguồn hoặc loại trừ dòng kèm lý do nghiệp vụ.",
        "BANK_PARSE_ERROR": "Sửa giá trị sai định dạng trong file nguồn hoặc loại trừ dòng kèm lý do nghiệp vụ.",
        "CURRENCY_MISMATCH": "Tách dữ liệu theo loại tiền hoặc sửa cột loại tiền đã ghép.",
    }
    details: list[dict[str, Any]] = []
    for issue_index, issue in enumerate(report.issues):
        table = tables.get(issue.feed)
        rows: list[dict[str, Any]] = []
        for row_number in issue.source_rows[:50]:
            raw_values: dict[str, Any] = {}
            if table is not None and 2 <= row_number <= len(table.frame) + 1:
                raw_values = {
                    str(key): str(value) if value is not None else ""
                    for key, value in table.frame.iloc[row_number - 2].items()
                }
            normalized_values = canonical_by_row.get((issue.feed, row_number))
            rows.append(
                {
                    "source_row_number": row_number,
                    "source_filename": file_meta.get("orders" if issue.feed == "orders" else "bank", {}).get("name"),
                    "sheet_name": table.sheet_name if table is not None else None,
                    "raw_values": raw_values,
                    "normalized_values": normalized_values,
                    "source_row_hash": (normalized_values or {}).get("lineage", {}).get("source_row_hash"),
                }
            )
        details.append(
            {
                "issue_index": issue_index,
                "code": issue.code,
                "severity": issue.severity.value,
                "feed": issue.feed,
                "message_vi": issue.message_vi,
                "sample_values": list(issue.sample_values),
                "suggested_action_vi": actions.get(
                    issue.code,
                    "Kiểm tra dữ liệu nguồn và cách ghép cột; việc xem chi tiết không thay đổi dữ liệu.",
                ),
                "rows": rows,
                "row_count": len(issue.source_rows),
                "is_truncated": len(issue.source_rows) > 50,
            }
        )
    return details
