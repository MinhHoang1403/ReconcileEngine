from pathlib import Path

import pytest
from openpyxl import Workbook

from smartrecon.ingest.readers import InputFileError, read_source_bundle, read_table


def test_semicolon_csv_with_intro_rows_keeps_exact_lineage(tmp_path: Path) -> None:
    source = tmp_path / "receivables.csv"
    source.write_text(
        "Báo cáo khoản phải thu;;;\n"
        "Ngày xuất: 16/10/2026;;;\n"
        "Mã ĐH;Ngày tạo;Thành tiền;Trạng thái\n"
        "00017;15/10/2026;1.250.000;Đã giao\n"
        "00018;16/10/2026;500.000;Đã giao\n",
        encoding="utf-8-sig",
    )

    table = read_table(source, expected_feed="orders")

    assert table.structure.status == "AUTO_SELECTED"
    assert table.structure.selected.delimiter == ";"
    assert table.structure.selected.header_row_number == 3
    assert table.source_row_numbers == (4, 5)
    assert table.frame.iloc[0]["Mã ĐH"] == "00017"


def test_multisheet_workbook_selects_only_business_sheet(tmp_path: Path) -> None:
    source = tmp_path / "bank.xlsx"
    workbook = Workbook()
    cover = workbook.active
    cover.title = "Giới thiệu"
    cover.append(["Báo cáo ngân hàng"])
    data = workbook.create_sheet("Giao dịch")
    data.append(["Mã giao dịch", "Ngày hạch toán", "Số tiền ghi có", "Nội dung CK"])
    data.append(["00001", "2026-10-15", 250000, "Thanh toán ORD-01"])
    workbook.save(source)
    workbook.close()

    table = read_table(source, expected_feed="bank_transactions")

    assert table.sheet_name == "Giao dịch"
    assert table.structure.status == "AUTO_SELECTED"
    assert table.source_row_numbers == (2,)
    assert table.frame.iloc[0]["Mã giao dịch"] == "00001"


def test_equally_plausible_sheets_require_confirmation(tmp_path: Path) -> None:
    source = tmp_path / "ambiguous.xlsx"
    workbook = Workbook()
    for index, title in enumerate(("Tháng 9", "Tháng 10")):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet.title = title
        sheet.append(["Mã ĐH", "Ngày tạo", "Thành tiền"])
        sheet.append([f"O-{index}", "2026-10-15", 100000])
    workbook.save(source)
    workbook.close()

    provisional = read_table(source, expected_feed="orders")
    confirmed = read_table(source, expected_feed="orders", sheet_name="Tháng 10", header_row=1)

    assert provisional.structure.status == "NEEDS_REVIEW"
    assert any(issue.code == "AMBIGUOUS_STRUCTURE" for issue in provisional.structure.issues)
    assert {candidate.sheet_name for candidate in provisional.structure.candidates[:2]} == {"Tháng 9", "Tháng 10"}
    assert confirmed.structure.status == "AUTO_SELECTED"
    assert confirmed.sheet_name == "Tháng 10"


def test_safe_two_row_header_is_flattened(tmp_path: Path) -> None:
    source = tmp_path / "split-bank.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Giao dịch", None, "Số tiền", None])
    sheet.append(["Mã", "Ngày", "Ghi có", "Ghi nợ"])
    sheet.append(["00001", "2026-10-15", 250000, None])
    workbook.save(source)
    workbook.close()

    table = read_table(source, expected_feed="bank_transactions")

    assert table.structure.selected.header_depth == 2
    assert table.source_row_numbers == (3,)
    assert table.frame.columns.tolist() == [
        "Giao dịch / Mã",
        "Giao dịch / Ngày",
        "Số tiền / Ghi có",
        "Số tiền / Ghi nợ",
    ]


def test_footer_total_is_retained_and_flagged(tmp_path: Path) -> None:
    source = tmp_path / "orders.csv"
    source.write_text(
        "order_id,ordered_at,expected_amount\nO1,2026-10-15,100000\nTổng cộng,,100000\n",
        encoding="utf-8",
    )

    table = read_table(source, expected_feed="orders")

    assert table.source_row_numbers == (2, 3)
    assert table.frame.iloc[1]["order_id"] == "Tổng cộng"
    assert table.structure.possible_footer_rows == (3,)
    assert any(issue.code == "POSSIBLE_FOOTER_ROWS" for issue in table.structure.issues)


def test_duplicate_header_is_rejected_instead_of_selecting_data_as_header(tmp_path: Path) -> None:
    source = tmp_path / "duplicate.csv"
    source.write_text("id,id\n1,2\n", encoding="utf-8")

    with pytest.raises(InputFileError, match="Header"):
        read_table(source)


def test_combined_workbook_auto_selects_distinct_feeds(tmp_path: Path) -> None:
    source = tmp_path / "combined.xlsx"
    workbook = Workbook()
    orders = workbook.active
    orders.title = "Đơn hàng"
    orders.append(["Mã ĐH", "Ngày tạo", "Thành tiền"])
    orders.append(["O1", "2026-10-15", 100000])
    bank = workbook.create_sheet("Ngân hàng")
    bank.append(["Mã giao dịch", "Ngày hạch toán", "Số tiền ghi có"])
    bank.append(["B1", "2026-10-16", 100000])
    workbook.save(source)
    workbook.close()

    bundle = read_source_bundle(workbook_path=source)

    assert bundle.orders.sheet_name == "Đơn hàng"
    assert bundle.bank_transactions.sheet_name == "Ngân hàng"
