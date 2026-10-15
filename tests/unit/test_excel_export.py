from pathlib import Path

from openpyxl import load_workbook

from backend.smartrecon_api.config import get_settings
from backend.smartrecon_api.services.exports import create_excel_export, safe_excel_value


def test_formula_injection_is_exported_as_text(tmp_path: Path, monkeypatch) -> None:
    settings = get_settings()
    monkeypatch.setattr(settings, "artifacts_path", tmp_path)
    result = {
        "kpis": {"currency": "VND", "confirmed_cash_allocated_minor": 28_000},
        "orders": [
            {
                "order_id": "O28",
                "description": '=HYPERLINK("https://invalid")',
                "expected_amount_minor": 28_000,
            }
        ],
        "bank_transactions": [],
        "groups": [],
        "allocations": [],
        "exceptions": [],
        "order_outstanding_minor": {"O28": 28_000},
        "payment_statuses": {"O28": "UNPAID"},
        "bank_unapplied_minor": {},
        "semantic_checksum": "fixture",
    }

    path = create_excel_export("formula-test", result)
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["02_Orders"]
    description_cell = sheet.cell(row=2, column=2)

    assert description_cell.data_type == "s"
    assert description_cell.value.startswith("'")
    assert safe_excel_value("-123") == "'-123"
    assert safe_excel_value(-123) == -123
