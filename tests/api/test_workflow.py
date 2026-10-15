from collections.abc import Generator
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.smartrecon_api.config import get_settings
from backend.smartrecon_api.db.database import get_db
from backend.smartrecon_api.db.models import ReconciliationRun, UploadSession
from backend.smartrecon_api.main import app


class MemorySession:
    """Small Session test double; PostgreSQL behavior is covered by Docker acceptance."""

    def __init__(self) -> None:
        self.objects: dict[tuple[type, str], object] = {}

    def add(self, value: object) -> None:
        identifier = getattr(value, "id", None)
        if isinstance(identifier, str):
            self.objects[(type(value), identifier)] = value

    def get(self, model: type, identifier: str):
        return self.objects.get((model, identifier))

    def flush(self) -> None:
        return None

    def commit(self) -> None:
        return None

    def refresh(self, value: object) -> None:
        return None

    def scalar(self, statement: object):
        return None


def test_demo_to_export_uses_one_persisted_run(tmp_path: Path, monkeypatch) -> None:
    session = MemorySession()

    def test_db() -> Generator[MemorySession, None, None]:
        yield session

    monkeypatch.setattr(get_settings(), "artifacts_path", tmp_path / "artifacts")
    app.dependency_overrides[get_db] = test_db
    try:
        with TestClient(app) as client:
            demo = client.post("/api/demo/datasets")
            assert demo.status_code == 200
            inspection = demo.json()
            assert "truth" not in demo.text.lower()

            mapping = {
                **inspection["suggested_mapping"],
                "currency": "VND",
                "source_profile": "BANK_CREDIT",
                "date_format": "%Y-%m-%d",
                "decimal_separator": ".",
                "thousand_separator": ",",
                "order_defaults": {"order_status": "ELIGIBLE"},
                "bank_defaults": {},
            }
            gate = client.post(f"/api/uploads/{inspection['session_id']}/mapping", json=mapping)
            assert gate.status_code == 200
            assert gate.json()["report"]["status"] == "READY"

            warning_codes = [
                issue["code"] for issue in gate.json()["report"]["issues"] if issue["severity"] == "WARNING"
            ]
            resolution = client.post(
                f"/api/uploads/{inspection['session_id']}/resolutions",
                json={"exclusions": [], "acknowledged_issue_codes": warning_codes},
            )
            assert resolution.status_code == 200
            assert resolution.json()["acknowledged_issue_codes"] == warning_codes

            created = client.post(
                "/api/runs",
                json={"upload_session_id": inspection["session_id"], "policy": {"currency": "VND"}},
            )
            assert created.status_code == 200
            run_id = created.json()["run_id"]
            repeated = client.post(
                "/api/runs",
                json={
                    "upload_session_id": inspection["session_id"],
                    "policy": {"currency": "VND", "fuzzy_candidate_threshold": 85},
                },
            )
            assert repeated.status_code == 200
            assert repeated.json()["semantic_checksum"] == created.json()["semantic_checksum"]
            dashboard = client.get(f"/api/runs/{run_id}/dashboard")
            assert dashboard.status_code == 200
            summary = dashboard.json()
            assert summary["conservation"]["is_valid"] is True
            assert summary["conservation"]["violations"] == []

            groups = client.get(f"/api/runs/{run_id}/groups?page_size=200")
            orders = client.get(f"/api/runs/{run_id}/orders?page_size=200")
            banks = client.get(f"/api/runs/{run_id}/bank-transactions?page_size=200")
            assert groups.json()["total"] > 0
            assert orders.json()["total"] == 100
            assert banks.json()["total"] == 98

            exported = client.get(f"/api/runs/{run_id}/exports?format=xlsx")
            assert exported.status_code == 200
            export_path = tmp_path / "reconciliation.xlsx"
            export_path.write_bytes(exported.content)
    finally:
        app.dependency_overrides.clear()

    assert session.get(UploadSession, inspection["session_id"]) is not None
    assert session.get(ReconciliationRun, run_id) is not None

    workbook = load_workbook(export_path, data_only=False, read_only=True)
    assert workbook.sheetnames == [
        "01_TongQuan_KPI",
        "02_Orders",
        "03_BankTransactions",
        "04_ConfirmedGroups",
        "05_AllocationEdges",
        "06_ReviewQueue",
        "07_Exceptions_Orphans",
        "08_ReviewAudit",
        "09_RunManifest",
    ]
    edge_sheet = workbook["05_AllocationEdges"]
    headers = [cell.value for cell in next(edge_sheet.iter_rows(max_row=1))]
    cash_column = headers.index("allocated_cash_minor") + 1
    exported_cash = sum(
        int(row[0] or 0)
        for row in edge_sheet.iter_rows(min_row=2, min_col=cash_column, max_col=cash_column, values_only=True)
    )
    assert exported_cash == summary["kpis"]["confirmed_cash_allocated_minor"]
