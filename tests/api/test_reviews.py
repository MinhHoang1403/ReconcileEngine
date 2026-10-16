from collections.abc import Generator
from io import BytesIO
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.smartrecon_api.config import get_settings
from backend.smartrecon_api.db.database import get_db
from backend.smartrecon_api.db.models import ReconciliationRun, ReviewAction
from backend.smartrecon_api.main import app
from backend.smartrecon_api.schemas.reviews import ApproveReviewRequest
from backend.smartrecon_api.services.reviews import approve_review
from backend.smartrecon_api.services.runs import (
    effective_result_path,
    load_run_result,
    result_from_payload,
    result_to_payload,
)
from backend.smartrecon_api.services.serialization import write_json
from smartrecon.config import ReconciliationPolicy
from smartrecon.domain.statuses import Currency
from smartrecon.engine import ReconciliationResult, run_reconciliation
from smartrecon.ingest.mapping import default_demo_mapping
from smartrecon.ingest.normalization import normalize_bundle
from smartrecon.ingest.readers import read_source_bundle

DEMO_DIR = Path("data/demo/vn_synthetic_demo")


class ReviewSession:
    """Transaction-aware test double; PostgreSQL locking is verified in Docker acceptance."""

    def __init__(self, run: ReconciliationRun, *, fail_commit: bool = False) -> None:
        self.run = run
        self.fail_commit = fail_commit
        self.committed_revision = run.current_revision
        self.committed_actions: dict[str, ReviewAction] = {}
        self.staged: list[object] = []
        self.rollback_count = 0

    def scalar(self, statement: object):
        entity = statement.column_descriptions[0].get("entity")
        if entity is ReconciliationRun:
            return self.run
        if entity is ReviewAction:
            params = statement.compile().params
            key = next(
                (value for name, value in params.items() if "idempotency_key" in name),
                None,
            )
            return next(
                (action for action in self.committed_actions.values() if action.idempotency_key == key),
                None,
            )
        if entity is None:
            return len(self.committed_actions)
        return None

    def scalars(self, statement: object):
        actions = sorted(
            self.committed_actions.values(),
            key=lambda item: item.revision_after,
            reverse=True,
        )

        class Rows:
            def all(self) -> list[ReviewAction]:
                return actions

        return Rows()

    def get(self, model: type, identifier: str):
        if model is ReviewAction:
            return self.committed_actions.get(identifier)
        if model is ReconciliationRun and identifier == self.run.id:
            return self.run
        return None

    def add(self, value: object) -> None:
        self.staged.append(value)

    def commit(self) -> None:
        if self.fail_commit:
            raise RuntimeError("forced commit failure")
        for value in self.staged:
            if isinstance(value, ReviewAction):
                self.committed_actions[value.id] = value
        self.staged.clear()
        self.committed_revision = self.run.current_revision

    def rollback(self) -> None:
        self.rollback_count += 1
        self.staged.clear()
        self.run.current_revision = self.committed_revision


@pytest.fixture(scope="module")
def base_result() -> ReconciliationResult:
    raw = read_source_bundle(
        orders_path=DEMO_DIR / "internal_receivables.csv",
        bank_path=DEMO_DIR / "bank_transactions.csv",
    )
    normalized = normalize_bundle(raw, default_demo_mapping())
    return run_reconciliation(normalized, ReconciliationPolicy(currency=Currency.VND))


def _group_id(result: ReconciliationResult) -> str:
    return next(
        group.group_id
        for group in result.groups
        if group.order_ids == ("ORD00009", "ORD00010") and group.bank_txn_ids == ("TXN00007",)
    )


def _proposal_id(
    result: ReconciliationResult,
    order_ids: tuple[str, ...],
    bank_ids: tuple[str, ...],
) -> str:
    return next(
        group.group_id for group in result.groups if group.order_ids == order_ids and group.bank_txn_ids == bank_ids
    )


def _session(tmp_path: Path, monkeypatch, result: ReconciliationResult, *, fail_commit: bool = False):
    monkeypatch.setattr(get_settings(), "artifacts_path", tmp_path / "artifacts")
    run = ReconciliationRun(
        id="run-review-test",
        upload_session_id="upload-review-test",
        status="COMPLETED",
        currency="VND",
        policy_snapshot={},
        mapping_snapshot={},
        semantic_checksum=result.semantic_checksum,
        current_revision=0,
    )
    session = ReviewSession(run, fail_commit=fail_commit)
    write_json(effective_result_path(run.id, 0), result_to_payload(result))
    return session


def test_review_api_is_idempotent_and_stale_revision_returns_409(
    tmp_path: Path,
    monkeypatch,
    base_result: ReconciliationResult,
) -> None:
    session = _session(tmp_path, monkeypatch, base_result)

    def test_db() -> Generator[ReviewSession, None, None]:
        yield session

    group_id = _group_id(base_result)
    competing_group_id = next(
        group.group_id
        for group in base_result.groups
        if group.group_id != group_id and "TXN00007" in group.bank_txn_ids
    )
    command = {
        "group_id": group_id,
        "expected_revision": 0,
        "idempotency_key": "approve-review-0001",
        "actor_label": "local_demo_label",
        "note": "Đã đối chiếu chứng từ nguồn",
    }
    app.dependency_overrides[get_db] = test_db
    try:
        with TestClient(app) as client:
            accepted = client.post(
                f"/api/runs/{session.run.id}/reviews/approve",
                json=command,
            )
            assert accepted.status_code == 200
            assert accepted.json()["effective_revision"] == 1
            assert accepted.json()["allocation_delta_minor"] == 500_000
            assert accepted.json()["conservation"]["is_valid"] is True

            replay = client.post(
                f"/api/runs/{session.run.id}/reviews/approve",
                json=command,
            )
            assert replay.status_code == 200
            assert replay.json()["action_id"] == accepted.json()["action_id"]
            assert replay.json()["idempotent_replay"] is True

            key_reuse = client.post(
                f"/api/runs/{session.run.id}/reviews/approve",
                json=command | {"group_id": "different-group"},
            )
            assert key_reuse.status_code == 409
            assert key_reuse.json()["detail"]["error_code"] == "IDEMPOTENCY_KEY_REUSED"

            stale = client.post(
                f"/api/runs/{session.run.id}/reviews/approve",
                json=command
                | {
                    "group_id": competing_group_id,
                    "idempotency_key": "approve-review-0002",
                },
            )
            assert stale.status_code == 409
            assert stale.json()["detail"]["error_code"] == "STALE_REVISION"
            assert stale.json()["detail"]["current_revision"] == 1
    finally:
        app.dependency_overrides.clear()

    assert len(session.committed_actions) == 1
    assert session.run.current_revision == 1


def test_review_commit_failure_rolls_back_and_removes_unpublished_artifact(
    tmp_path: Path,
    monkeypatch,
    base_result: ReconciliationResult,
) -> None:
    session = _session(tmp_path, monkeypatch, base_result, fail_commit=True)
    request = ApproveReviewRequest(
        group_id=_group_id(base_result),
        expected_revision=0,
        idempotency_key="approve-review-failure",
        note="Đã đối chiếu chứng từ nguồn",
    )

    with pytest.raises(RuntimeError, match="forced commit failure"):
        approve_review(session, session.run.id, request)

    assert session.rollback_count == 1
    assert session.run.current_revision == 0
    assert session.committed_actions == {}
    assert not effective_result_path(session.run.id, 1).exists()


def test_undo_creates_a_new_revision_and_restores_the_prior_snapshot(
    tmp_path: Path,
    monkeypatch,
    base_result: ReconciliationResult,
) -> None:
    session = _session(tmp_path, monkeypatch, base_result)

    def test_db() -> Generator[ReviewSession, None, None]:
        yield session

    app.dependency_overrides[get_db] = test_db
    try:
        with TestClient(app) as client:
            approved = client.post(
                f"/api/runs/{session.run.id}/reviews/approve",
                json={
                    "group_id": _group_id(base_result),
                    "expected_revision": 0,
                    "idempotency_key": "undo-approve-review-0001",
                    "note": "Đã đối chiếu chứng từ nguồn",
                },
            )
            queue_after_approval = client.get(f"/api/runs/{session.run.id}/reviews?page_size=20")
            history_after_approval = client.get(f"/api/runs/{session.run.id}/reviews/history?page_size=20")
            dashboard_after_approval = client.get(f"/api/runs/{session.run.id}/dashboard")
            export_after_approval = client.get(f"/api/runs/{session.run.id}/exports?format=xlsx&expected_revision=1")
            undone = client.post(
                f"/api/runs/{session.run.id}/reviews/undo",
                json={
                    "action_id": approved.json()["action_id"],
                    "expected_revision": 1,
                    "idempotency_key": "undo-review-action-0001",
                    "note": "Khôi phục proposal để kiểm tra lại",
                },
            )
            queue_after_undo = client.get(f"/api/runs/{session.run.id}/reviews")
            dashboard_after_undo = client.get(f"/api/runs/{session.run.id}/dashboard")
            historical_dashboard = client.get(f"/api/runs/{session.run.id}/dashboard?revision=1")
            historical_groups = client.get(f"/api/runs/{session.run.id}/groups?revision=1&page_size=200")
            stale_export = client.get(f"/api/runs/{session.run.id}/exports?format=xlsx&expected_revision=1")
    finally:
        app.dependency_overrides.clear()

    assert approved.status_code == 200
    assert queue_after_approval.status_code == 200
    assert queue_after_approval.json()["effective_revision"] == 1
    assert _group_id(base_result) not in {item["group_id"] for item in queue_after_approval.json()["items"]}
    assert history_after_approval.json()["total"] == 1
    assert history_after_approval.json()["items"][0]["actor_is_authenticated"] is False
    assert dashboard_after_approval.status_code == 200
    assert dashboard_after_approval.json()["effective_revision"] == 1
    assert dashboard_after_approval.json()["analytics"]["review_status_counts"]["APPROVED"] == 1
    assert dashboard_after_approval.json()["analytics"]["review_decision_counts"]["APPROVE_GROUP"] == 1
    assert export_after_approval.status_code == 200
    assert undone.status_code == 200
    assert undone.json()["effective_revision"] == 2
    assert undone.json()["allocation_delta_minor"] == -500_000
    assert undone.json()["conservation"]["is_valid"] is True
    assert session.run.current_revision == 2
    assert len(session.committed_actions) == 2
    restored = result_from_payload(load_run_result(session.run.id, 2))
    restored_group = next(group for group in restored.groups if group.group_id == _group_id(base_result))
    assert restored_group.link_status.value == "PROPOSED"
    assert restored_group.review_status.value == "UNDONE"

    restored_item = next(
        item for item in queue_after_undo.json()["items"] if item["group_id"] == _group_id(base_result)
    )
    assert restored_item["review_status"] == "UNDONE"
    assert restored_item["can_approve"] is True
    assert dashboard_after_undo.json()["effective_revision"] == 2
    assert dashboard_after_undo.json()["analytics"]["review_status_counts"]["UNDONE"] == 1
    assert dashboard_after_undo.json()["analytics"]["review_decision_counts"]["UNDO_ACTION"] == 1
    assert historical_dashboard.json()["effective_revision"] == 1
    assert historical_dashboard.json()["kpis"] == dashboard_after_approval.json()["kpis"]
    assert historical_groups.json()["effective_revision"] == 1
    historical_target = next(
        item for item in historical_groups.json()["items"] if item["group_id"] == _group_id(base_result)
    )
    assert historical_target["review_status"] == "APPROVED"
    assert stale_export.status_code == 409
    assert stale_export.json()["detail"]["error_code"] == "STALE_REVISION"
    assert restored.allocations == base_result.allocations
    assert restored.kpis == base_result.kpis
    assert restored.conservation == base_result.conservation

    workbook = load_workbook(BytesIO(export_after_approval.content), data_only=False)
    assert len(workbook.sheetnames) == 9
    audit = workbook["08_ReviewAudit"]
    audit_headers = [cell.value for cell in audit[1]]
    assert audit.max_row == 2
    assert audit.cell(row=2, column=audit_headers.index("Hành động") + 1).value == "Chấp thuận đề xuất"
    assert audit.cell(row=2, column=audit_headers.index("Revision sau") + 1).value == 1
    confirmed = workbook["04_ConfirmedGroups"]
    confirmed_headers = [cell.value for cell in confirmed[1]]
    approved_row = next(
        row
        for row in confirmed.iter_rows(min_row=2)
        if row[confirmed_headers.index("Mã nhóm")].value == _group_id(base_result)
    )
    assert approved_row[confirmed_headers.index("Trạng thái xác minh")].value == "Đã chấp thuận"
    allocation = workbook["05_AllocationEdges"]
    allocation_headers = [cell.value for cell in allocation[1]]
    cash_column = allocation_headers.index("Số tiền phân bổ")
    assert (
        sum(row[cash_column].value for row in allocation.iter_rows(min_row=2))
        == dashboard_after_approval.json()["kpis"]["confirmed_cash_allocated_minor"]
    )


def test_undo_is_blocked_when_a_downstream_review_revision_exists(
    tmp_path: Path,
    monkeypatch,
    base_result: ReconciliationResult,
) -> None:
    session = _session(tmp_path, monkeypatch, base_result)

    def test_db() -> Generator[ReviewSession, None, None]:
        yield session

    app.dependency_overrides[get_db] = test_db
    try:
        with TestClient(app) as client:
            approved = client.post(
                f"/api/runs/{session.run.id}/reviews/approve",
                json={
                    "group_id": _group_id(base_result),
                    "expected_revision": 0,
                    "idempotency_key": "dependency-approve-0001",
                    "note": "Đã đối chiếu chứng từ nguồn",
                },
            )
            rejected = client.post(
                f"/api/runs/{session.run.id}/reviews/reject",
                json={
                    "group_id": _proposal_id(base_result, ("ORD00004",), ("TXN00004",)),
                    "expected_revision": 1,
                    "idempotency_key": "dependency-reject-0001",
                    "note": "Tham chiếu nguồn không phù hợp",
                },
            )
            blocked = client.post(
                f"/api/runs/{session.run.id}/reviews/undo",
                json={
                    "action_id": approved.json()["action_id"],
                    "expected_revision": 2,
                    "idempotency_key": "dependency-undo-0001",
                },
            )
    finally:
        app.dependency_overrides.clear()

    assert approved.status_code == 200
    assert rejected.status_code == 200
    assert blocked.status_code == 409
    assert blocked.json()["detail"]["error_code"] == "UNDO_DEPENDENCY_CONFLICT"
    assert session.run.current_revision == 2
    assert len(session.committed_actions) == 2
    assert not effective_result_path(session.run.id, 3).exists()
