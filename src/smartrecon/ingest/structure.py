"""Deterministic CSV/XLSX structure inspection with conservative ambiguity flags."""

from __future__ import annotations

import csv
import re
import unicodedata
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

COMMON_DELIMITERS = (",", ";", "\t", "|")
MAX_HEADER_SCAN_ROWS = 30
MAX_STRUCTURE_PREVIEW_ROWS = 80

_FEED_TERMS = {
    "orders": {
        "order",
        "invoice",
        "receivable",
        "amount_due",
        "expected_amount",
        "ma_don_hang",
        "so_don_hang",
        "ngay_tao",
        "ngay_dat_hang",
        "thanh_tien",
        "so_tien_can_thu",
    },
    "bank_transactions": {
        "transaction",
        "bank",
        "posting_date",
        "credit_amount",
        "debit_amount",
        "posted_amount",
        "ma_giao_dich",
        "ngay_hach_toan",
        "so_tien_ghi_co",
        "noi_dung_ck",
    },
}

_TOTAL_LABELS = {
    "total",
    "grand_total",
    "subtotal",
    "tong",
    "tong_cong",
    "cong",
}


def normalize_structure_text(value: object) -> str:
    """Normalize labels for structure scoring without changing source values."""
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value).strip().casefold())
    text = "".join(character for character in text if not unicodedata.combining(character))
    text = text.replace("đ", "d")
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


@dataclass(frozen=True, slots=True)
class StructureCandidate:
    """One plausible sheet/header interpretation ranked by deterministic rules."""

    sheet_name: str | None
    header_row_number: int
    header_depth: int
    delimiter: str | None
    score: int
    reasons: tuple[str, ...]
    columns: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class StructureIssue:
    """A structural condition that must stay visible to the user."""

    code: str
    status: str
    message_vi: str
    source_rows: tuple[int, ...] = ()


@dataclass(frozen=True, slots=True)
class DetectedStructure:
    """Selected interpretation plus competing candidates and audit metadata."""

    status: str
    selected: StructureCandidate
    candidates: tuple[StructureCandidate, ...]
    issues: tuple[StructureIssue, ...]
    encoding: str | None
    data_start_row_number: int
    source_row_numbers: tuple[int, ...]
    ignored_blank_rows: tuple[int, ...]
    possible_footer_rows: tuple[int, ...]


def _nonempty(row: Iterable[object]) -> list[object]:
    return [value for value in row if value is not None and str(value).strip()]


def _looks_like_data_value(value: object) -> bool:
    if value is None or not str(value).strip():
        return False
    text = str(value).strip()
    if re.fullmatch(r"[-+]?\d[\d., ]*", text):
        return True
    return bool(
        re.fullmatch(r"\d{4}[-/]\d{1,2}[-/]\d{1,2}(?:[ T].*)?", text)
        or re.fullmatch(r"\d{1,2}[-/]\d{1,2}[-/]\d{2,4}(?:[ T].*)?", text)
    )


def _flatten_headers(rows: list[list[object]], row_index: int, depth: int) -> tuple[str, ...]:
    width = max((len(row) for row in rows[row_index : row_index + depth]), default=0)
    headers: list[str] = []
    carried_top = ""
    for column_index in range(width):
        parts: list[str] = []
        for depth_index in range(depth):
            row = rows[row_index + depth_index]
            value = row[column_index] if column_index < len(row) else None
            text = "" if value is None else str(value).strip()
            if depth == 2 and depth_index == 0:
                if text:
                    carried_top = text
                elif carried_top:
                    text = carried_top
            if text and text not in parts:
                parts.append(text)
        headers.append(" / ".join(parts))
    while headers and not headers[-1]:
        headers.pop()
    return tuple(headers)


def _header_score(
    rows: list[list[object]],
    row_index: int,
    depth: int,
    expected_feed: str | None,
) -> tuple[int, tuple[str, ...], tuple[str, ...]]:
    if row_index + depth >= len(rows):
        return -1, (), ()
    if depth == 2:
        first_level = [normalize_structure_text(value) for value in rows[row_index]]
        populated_first = [value for value in first_level if value]
        if len(populated_first) < 2:
            return -1, (), ()
        if (
            populated_first
            and len(populated_first) == len(first_level)
            and len(set(populated_first)) == len(populated_first)
        ):
            return -1, (), ()
        second_level = _nonempty(rows[row_index + 1])
        if second_level and sum(_looks_like_data_value(value) for value in second_level) / len(second_level) > 0.35:
            return -1, (), ()
    headers = _flatten_headers(rows, row_index, depth)
    populated = [header for header in headers if header]
    if len(populated) < 2:
        return -1, (), headers
    normalized = [normalize_structure_text(header) for header in populated]
    if len(set(normalized)) != len(normalized):
        return -1, ("Header trùng sau khi chuẩn hóa",), headers

    text_headers = sum(not _looks_like_data_value(header) for header in populated)
    name_score = round(25 * text_headers / len(populated))
    unique_score = round(15 * len(set(normalized)) / len(normalized))
    semantic_terms = _FEED_TERMS.get(expected_feed or "", set())
    semantic_hits = sum(
        any(term == label or term in label.split("_") or term in label for term in semantic_terms)
        for label in normalized
    )
    semantic_score = min(35, semantic_hits * 9)

    data_rows = [row for row in rows[row_index + depth : row_index + depth + 8] if _nonempty(row)]
    if data_rows:
        expected_width = len(headers)
        consistent = sum(abs(len(row) - expected_width) <= 1 for row in data_rows)
        density = sum(min(len(_nonempty(row)), expected_width) for row in data_rows) / (
            len(data_rows) * max(expected_width, 1)
        )
        shape_score = round(15 * consistent / len(data_rows)) + round(10 * density)
    else:
        shape_score = 0

    score = name_score + unique_score + semantic_score + shape_score
    reasons = [
        f"{len(populated)} tên cột không trống và không trùng",
        f"{len(data_rows)} dòng kế tiếp có cấu trúc dữ liệu",
    ]
    if semantic_hits:
        reasons.append(f"{semantic_hits} tên cột phù hợp ngữ nghĩa nguồn")
    if depth == 2:
        reasons.append("Ghép an toàn hai hàng tiêu đề")
    return score, tuple(reasons), headers


def _rank_candidates(
    sheets: dict[str | None, list[list[object]]],
    *,
    expected_feed: str | None,
    delimiter: str | None,
) -> list[StructureCandidate]:
    candidates: list[StructureCandidate] = []
    for sheet_name, rows in sheets.items():
        scan_limit = min(MAX_HEADER_SCAN_ROWS, max(len(rows) - 1, 0))
        for row_index in range(scan_limit):
            for depth in (1, 2):
                score, reasons, columns = _header_score(rows, row_index, depth, expected_feed)
                if score < 0:
                    continue
                candidates.append(
                    StructureCandidate(
                        sheet_name=sheet_name,
                        header_row_number=row_index + 1,
                        header_depth=depth,
                        delimiter=delimiter,
                        score=score,
                        reasons=reasons,
                        columns=columns,
                    )
                )
    return sorted(
        candidates,
        key=lambda candidate: (
            -candidate.score,
            candidate.header_row_number,
            candidate.header_depth,
            candidate.sheet_name or "",
        ),
    )


def _detect_footer_rows(rows: list[list[object]], source_rows: list[int]) -> tuple[int, ...]:
    footer_rows: list[int] = []
    for row, source_row in zip(rows[-10:], source_rows[-10:], strict=False):
        values = _nonempty(row)
        if not values:
            continue
        first = normalize_structure_text(values[0])
        if first in _TOTAL_LABELS or any(first.startswith(f"{label}_") for label in _TOTAL_LABELS):
            footer_rows.append(source_row)
    return tuple(footer_rows)


def detect_structure(
    sheets: dict[str | None, list[list[object]]],
    *,
    expected_feed: str | None,
    delimiter: str | None = None,
    encoding: str | None = None,
    selected_sheet: str | None = None,
    selected_header_row: int | None = None,
    selected_header_depth: int | None = None,
) -> DetectedStructure:
    """Rank possible structures and return a provisional, auditable selection."""
    candidates = _rank_candidates(sheets, expected_feed=expected_feed, delimiter=delimiter)
    if not candidates:
        for rows in sheets.values():
            for row in rows[:MAX_HEADER_SCAN_ROWS]:
                labels = [normalize_structure_text(value) for value in _nonempty(row)]
                if len(labels) >= 2 and len(set(labels)) != len(labels):
                    raise ValueError("Header bị trùng và không thể diễn giải an toàn")
        raise ValueError("Không tìm thấy hàng tiêu đề và vùng dữ liệu có thể diễn giải an toàn")

    explicit = selected_header_row is not None or selected_sheet is not None or selected_header_depth is not None
    filtered = [
        candidate
        for candidate in candidates
        if (selected_sheet is None or candidate.sheet_name == selected_sheet)
        and (selected_header_row is None or candidate.header_row_number == selected_header_row)
        and (selected_header_depth is None or candidate.header_depth == selected_header_depth)
    ]
    if not filtered:
        raise ValueError("Cấu hình sheet/header đã chọn không khớp cấu trúc file")
    selected = filtered[0]
    runner_up = next(
        (
            candidate
            for candidate in candidates
            if candidate.sheet_name != selected.sheet_name
            or candidate.header_row_number != selected.header_row_number
            or candidate.header_depth != selected.header_depth
        ),
        None,
    )
    margin = selected.score - (runner_up.score if runner_up else 0)
    status = "AUTO_SELECTED" if explicit or (selected.score >= 55 and margin >= 8) else "NEEDS_REVIEW"

    selected_rows = sheets[selected.sheet_name]
    data_start = selected.header_row_number + selected.header_depth
    raw_data_rows = selected_rows[data_start - 1 :]
    raw_source_rows = list(range(data_start, data_start + len(raw_data_rows)))
    ignored_blank_rows = tuple(
        source_row for row, source_row in zip(raw_data_rows, raw_source_rows, strict=False) if not _nonempty(row)
    )
    data_pairs = [
        (row, source_row) for row, source_row in zip(raw_data_rows, raw_source_rows, strict=False) if _nonempty(row)
    ]
    data_rows = [pair[0] for pair in data_pairs]
    source_rows = [pair[1] for pair in data_pairs]
    possible_footers = _detect_footer_rows(data_rows, source_rows)

    issues: list[StructureIssue] = []
    unnamed_positions = [index + 1 for index, header in enumerate(selected.columns) if not header]
    if unnamed_positions:
        issues.append(
            StructureIssue(
                "UNNAMED_COLUMNS",
                "NEEDS_REVIEW",
                "Có cột chưa có tên rõ ràng tại vị trí: " + ", ".join(map(str, unnamed_positions)),
            )
        )
    if status == "NEEDS_REVIEW":
        issues.append(
            StructureIssue(
                "AMBIGUOUS_STRUCTURE",
                "NEEDS_REVIEW",
                "Có nhiều cách chọn sheet hoặc hàng tiêu đề hợp lý; cần xác nhận trước khi chuẩn hóa.",
            )
        )
    if possible_footers:
        issues.append(
            StructureIssue(
                "POSSIBLE_FOOTER_ROWS",
                "NEEDS_REVIEW",
                "Phát hiện dòng có vẻ là tổng cộng; hệ thống giữ nguyên và chưa coi là giao dịch hợp lệ.",
                possible_footers,
            )
        )
    return DetectedStructure(
        status=status,
        selected=selected,
        candidates=tuple(candidates[:8]),
        issues=tuple(issues),
        encoding=encoding,
        data_start_row_number=data_start,
        source_row_numbers=tuple(source_rows),
        ignored_blank_rows=ignored_blank_rows,
        possible_footer_rows=possible_footers,
    )


def decode_csv(path: Path) -> tuple[str, str]:
    """Decode only accepted UTF-8 variants and report the effective label."""
    content = path.read_bytes()
    if b"\x00" in content[:4096]:
        raise UnicodeError("CSV chứa byte NUL hoặc content không phải text")
    encoding = "utf-8-sig" if content.startswith(b"\xef\xbb\xbf") else "utf-8"
    return content.decode(encoding), encoding


def sniff_delimiter(text: str) -> str:
    """Detect a common delimiter, rejecting ties instead of guessing."""
    sample = text[:65536]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="".join(COMMON_DELIMITERS))
        return dialect.delimiter
    except csv.Error as exc:
        nonempty_lines = [line for line in sample.splitlines() if line.strip()][:10]
        counts = {delimiter: sum(line.count(delimiter) for line in nonempty_lines) for delimiter in COMMON_DELIMITERS}
        best = max(counts.values(), default=0)
        winners = [delimiter for delimiter, count in counts.items() if count == best and count > 0]
        if len(winners) != 1:
            raise ValueError("Không xác định được dấu phân cách CSV một cách an toàn") from exc
        return winners[0]


def csv_rows(path: Path, delimiter: str | None = None) -> tuple[list[list[str]], str, str]:
    text, encoding = decode_csv(path)
    selected_delimiter = delimiter or sniff_delimiter(text)
    rows = list(csv.reader(text.splitlines(), delimiter=selected_delimiter))
    return rows, selected_delimiter, encoding


def _excel_cell_value(cell: Any) -> object:
    """Preserve displayed leading zeroes when Excel stores an ID as a number."""
    value = cell.value
    if isinstance(value, int) and cell.number_format and re.fullmatch(r"0+", cell.number_format):
        return f"{value:0{len(cell.number_format)}d}"
    return value


def xlsx_rows(path: Path) -> dict[str, list[list[object]]]:
    workbook = load_workbook(path, read_only=False, data_only=True, keep_links=False)
    try:
        result: dict[str, list[list[object]]] = {}
        for worksheet in workbook.worksheets:
            if worksheet.max_row > 100_030:
                raise ValueError("Workbook vượt giới hạn 100.000 dòng dữ liệu và vùng tiêu đề an toàn")
            if worksheet.max_column > 200:
                raise ValueError("Workbook vượt giới hạn 200 cột")
            rows: list[list[object]] = []
            for row in worksheet.iter_rows():
                rows.append([_excel_cell_value(cell) for cell in row])
            result[worksheet.title] = rows
        return result
    finally:
        workbook.close()
